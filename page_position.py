# -*- coding: utf-8 -*-
"""位置精度页（模板第 5 页起）：每个 profile 一页。

页面内容：标题 + 居中位置精度图（25×15cm）+ 图表下方结论。
"""
import os

import numpy as np
from pptx.util import Cm

import parameter as P
import pptx_utils as U

TITLE_PREFIX = "位置精度测量"
PLACEHOLDER_TEXT = "箱线图展示全部数据"
CONCL_PREFIX = "结论"

FONT_EA = P.FONT_EA_PPT
FONT_LATIN = P.FONT_LATIN_PPT
FONT_SIZE = P.FONT_SIZE_PPT_BODY


def get_profiles(data):
    """位置精度数据中的 profile 列表（排除统计）。"""
    return [k for k in data.get("位置精度", {}).keys() if k != "统计"]


def count_pages(data):
    """本模块会生成多少页。"""
    return len(get_profiles(data))


def profile_all_pass(data, profile_name):
    """判断该 profile 下所有器件的全部数据是否满足 SPEC。

    判据：各样本减均值后的最大偏差 <= 该器件位置 SPEC（排除 angle）。
    """
    spec = data.get("物料信息", {}).get("spec", {})
    group = data.get("位置精度", {}).get(profile_name, {})
    checked = False
    for dev_name, dev_data in group.items():
        if dev_name == "统计" or not isinstance(dev_data, dict):
            continue
        bound = P.parse_position_bound(spec.get(dev_name, {}).get("位置"))
        if bound is None:
            return False
        for dim, raw in dev_data.items():
            if dim == "统计" or str(dim).lower() == "angle":
                continue
            if not isinstance(raw, list):
                continue
            vals = [float(v) for v in raw if v is not None]
            if not vals:
                continue
            checked = True
            m = float(np.mean(vals))
            if max(abs(v - m) for v in vals) > bound:
                return False
    return checked


def _remove_placeholder(slide):
    shp = U.find_shape_containing(slide, PLACEHOLDER_TEXT)
    if shp is not None:
        shp._element.getparent().remove(shp._element)


def _insert_chart(slide, profile_name):
    img = os.path.join(P.POSITION_IMG_DIR, f"{profile_name}_position.png")
    if not os.path.exists(img):
        print(f"  [POS] 图片不存在: {img}")
        return False
    slide.shapes.add_picture(
        img, Cm(P.POSITION_IMG_LEFT_CM), Cm(P.POSITION_IMG_TOP_CM),
        width=Cm(P.POSITION_IMG_W_CM), height=Cm(P.POSITION_IMG_H_CM))
    print(f"  [POS] {profile_name} -> {os.path.basename(img)} "
          f"({P.POSITION_IMG_W_CM}x{P.POSITION_IMG_H_CM}cm)")
    return True


def _set_conclusion(slide, profile_name, all_pass):
    if all_pass:
        text = f"结论：{profile_name} 下，所有器件位置精度均满足SPEC"
    else:
        text = f"结论：{profile_name} 下，存在器件位置精度不满足SPEC"

    target = U.find_text_shape(slide, CONCL_PREFIX)
    if target is None:
        # 模板页无结论文本框，图片下方新建
        l, t, w, h = P.POSITION_CONCL_NEW_BOX
        target = slide.shapes.add_textbox(Cm(l), Cm(t), Cm(w), Cm(h))
    else:
        left, width = P.POSITION_CONCL_FIT_BOX
        target.left = Cm(left)
        target.width = Cm(width)
    U.write_text_box(target, text, FONT_LATIN, FONT_EA, FONT_SIZE)
    return text


def fill_profile_page(slide, profile_name, data):
    _remove_placeholder(slide)
    if not U.set_title(slide, TITLE_PREFIX, f"{TITLE_PREFIX}-{profile_name}"):
        print(f"  [POS] 未找到标题文本框: {profile_name}")
    _insert_chart(slide, profile_name)
    all_pass = profile_all_pass(data, profile_name)
    _set_conclusion(slide, profile_name, all_pass)
    print(f"  [POS] {profile_name} 结论: "
          f"{'全部满足SPEC' if all_pass else '存在不满足'}")
    return all_pass


def fill_position_pages(prs, data, template_index=None):
    """填充位置精度页：每个 profile 一页。

    template_index 由主流程规划后传入；页面复制已由主流程完成，
    本函数只往 [template_index, template_index+count) 里填内容。
    返回填充的页数。
    """
    if template_index is None:
        template_index = P.TEMPLATE_PAGE_INDEX["position"]

    profiles = get_profiles(data)
    if not profiles:
        print("[POS] 无位置精度 profile")
        return 0

    for i, prof in enumerate(profiles):
        fill_profile_page(prs.slides[template_index + i], prof, data)

    print(f"[POS] 共生成 {len(profiles)} 页位置精度: {profiles}")
    return len(profiles)
