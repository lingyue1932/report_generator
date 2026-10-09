# -*- coding: utf-8 -*-
"""位置精度页：每个profile生成一页，居中插入位置精度图 + 结论"""
import os
import copy

import numpy as np
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

import parameter as P

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

TEMPLATE_SLIDE_INDEX = 4  # 模板第五页（0-based）

FONT_EA = "微软雅黑"
FONT_LATIN = "Times New Roman"
FONT_SIZE = 14

# 图片固定尺寸（长*高，cm）与版面坐标：统一在 parameter.py 配置

PLACEHOLDER_TEXT = "箱线图展示全部数据"
TITLE_PREFIX = "位置精度测量"


def _parse_bound(pos_str):
    """从 '±0.02' 解析精度SPEC数值"""
    import re
    m = re.search(r'[±+-]?([\d.]+)', str(pos_str))
    return float(m.group(1)) if m else None


def profile_all_pass(data, profile_name):
    """判断该profile下所有器件的全部数据是否满足SPEC
    判据：各样本减均值后的最大偏差 <= 该器件位置SPEC（排除angle）
    """
    spec = data.get("物料信息", {}).get("spec", {})
    group = data.get("位置精度", {}).get(profile_name, {})
    checked = False
    for dev_name, dev_data in group.items():
        if dev_name == "统计" or not isinstance(dev_data, dict):
            continue
        bound = _parse_bound(spec.get(dev_name, {}).get("位置"))
        if bound is None:
            return False
        for dim, raw in dev_data.items():
            if dim == "统计" or str(dim).lower() == "angle":
                continue
            if not isinstance(raw, list):
                continue
            vals = [float(v) for v in raw if v is not None]
            if not vals:
                continue
            checked = True
            m = float(np.mean(vals))
            if max(abs(v - m) for v in vals) > bound:
                return False
    return checked


def _set_run_font(run, latin_font=FONT_LATIN, ea_font=FONT_EA,
                  size_pt=FONT_SIZE):
    run.font.name = latin_font
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn('a:latin'))
    if latin is None:
        latin = rPr.makeelement(qn('a:latin'), {'typeface': latin_font})
        rPr.append(latin)
    else:
        latin.set('typeface', latin_font)
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {'typeface': ea_font})
        latin.addnext(ea)
    else:
        ea.set('typeface', ea_font)
    cs = rPr.find(qn('a:cs'))
    if cs is None:
        cs = rPr.makeelement(qn('a:cs'), {'typeface': latin_font})
        ea.addnext(cs)
    else:
        cs.set('typeface', latin_font)


def _duplicate_slide(prs, src_index):
    """复制幻灯片（纯形状/文本，无图片关系），返回新slide（追加在末尾）"""
    src = prs.slides[src_index]
    dest = prs.slides.add_slide(src.slide_layout)
    # 移除layout带入的占位符
    for shp in list(dest.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in src.shapes:
        dest.shapes._spTree.append(copy.deepcopy(shp._element))
    return dest


def _reorder_slide(prs, from_index, to_index):
    """调整幻灯片顺序"""
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    el = ids[from_index]
    sldIdLst.remove(el)
    sldIdLst.insert(to_index, el)


def _remove_placeholder(slide):
    for shp in list(slide.shapes):
        if shp.has_text_frame and PLACEHOLDER_TEXT in shp.text_frame.text:
            shp._element.getparent().remove(shp._element)


def _insert_image(slide, profile_name):
    img = os.path.join(P.POSITION_IMG_DIR, f"{profile_name}_position.png")
    if not os.path.exists(img):
        print(f"  [POS] 图片不存在: {img}")
        return False
    slide.shapes.add_picture(
        img, Cm(P.POSITION_IMG_LEFT_CM), Cm(P.POSITION_IMG_TOP_CM),
        width=Cm(P.POSITION_IMG_W_CM), height=Cm(P.POSITION_IMG_H_CM))
    print(f"  [POS] {profile_name} -> {os.path.basename(img)} "
          f"({P.POSITION_IMG_W_CM}x{P.POSITION_IMG_H_CM}cm)")
    return True


def _set_conclusion(slide, profile_name, all_pass):
    if all_pass:
        text = f"结论：{profile_name} 下，所有器件位置精度均满足SPEC"
    else:
        text = f"结论：{profile_name} 下，存在器件位置精度不满足SPEC"

    target = None
    for shp in slide.shapes:
        if shp.has_text_frame and shp.text_frame.text.strip().startswith("结论"):
            target = shp
            break
    if target is None:
        # 模板页无结论文本框，图片下方新建
        l, t, w, h = P.POSITION_CONCL_NEW_BOX
        target = slide.shapes.add_textbox(Cm(l), Cm(t), Cm(w), Cm(h))
    else:
        left, width = P.POSITION_CONCL_FIT_BOX
        target.left = Cm(left)
        target.width = Cm(width)
    tf = target.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    _set_run_font(run)
    return text


def _set_title(slide, profile_name):
    """标题：位置精度测量-当前profile"""
    for shp in slide.shapes:
        if not shp.has_text_frame:
            continue
        if not shp.text_frame.text.strip().startswith(TITLE_PREFIX):
            continue
        tf = shp.text_frame
        title = f"{TITLE_PREFIX}-{profile_name}"
        p = tf.paragraphs[0]
        if p.runs:
            p.runs[0].text = title
            for r in p.runs[1:]:
                r.text = ""
        else:
            run = p.add_run()
            run.text = title
        for extra in tf.paragraphs[1:]:
            for r in extra.runs:
                r.text = ""
        return True
    print(f"  [POS] 未找到标题文本框: {profile_name}")
    return False


def _fill_profile_page(prs, slide, profile_name, data):
    _remove_placeholder(slide)
    _set_title(slide, profile_name)
    _insert_image(slide, profile_name)
    all_pass = profile_all_pass(data, profile_name)
    _set_conclusion(slide, profile_name, all_pass)
    print(f"  [POS] {profile_name} 结论: "
          f"{'全部满足SPEC' if all_pass else '存在不满足'}")


def get_profiles(data):
    """位置精度数据中的profile列表（排除统计）"""
    pos = data.get("位置精度", {})
    return [k for k in pos.keys() if k != "统计"]


def fill_position_pages(prs, data):
    """生成位置精度页：每个profile一页
    - 模板第五页作为第一个profile的页面，其余复制模板页插入其后
    - 图片居中，固定 25 x 15 cm
    - 图表下方写结论
    返回生成的页数
    """
    profiles = get_profiles(data)
    if not profiles:
        print("[POS] 无位置精度profile")
        return 0

    # 复制模板页（先复制再填充，保证副本是干净模板）
    for i in range(1, len(profiles)):
        _duplicate_slide(prs, TEMPLATE_SLIDE_INDEX)
        # 新页在末尾，移到模板页之后
        _reorder_slide(prs, len(prs.slides._sldIdLst) - 1,
                       TEMPLATE_SLIDE_INDEX + i)

    # 逐页填充
    for i, prof in enumerate(profiles):
        slide = prs.slides[TEMPLATE_SLIDE_INDEX + i]
        _fill_profile_page(prs, slide, prof, data)

    print(f"[POS] 共生成 {len(profiles)} 页位置精度: {profiles}")
    return len(profiles)
