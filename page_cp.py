# -*- coding: utf-8 -*-
"""CP 切片页（模板第 8 页起）：每个 <器件>CP 文件夹一页。

每页：标题 + 表格（一行一个 profile：条件 / 切片图 / 空洞率实测 / SPEC / 单行结论）
      + 整页结论。

切片图约定：<器件>CP/CP/<profile号>-<序号>/A.*，取名为 A 的图，
等比缩放至高度 parameter.CP_IMG_HEIGHT_CM。
"""
import os
import re

from pptx.util import Cm

import parameter as P
import pptx_utils as U

TITLE_PREFIX = "切片"
CONCL_PREFIX = "结论"

FONT_EA = P.FONT_EA_PPT
FONT_LATIN = P.FONT_LATIN_PPT
FONT_SIZE = P.FONT_SIZE_PPT_BODY

# 表格列
COL_COND, COL_IMG, COL_HOLE, COL_SPEC, COL_CONCL = 0, 1, 2, 3, 4

IMG_EXTS = (".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp")


# ============================================================
# 器件 / 图片发现
# ============================================================
def get_devices(base_dir=None):
    """扫描 xxxCP 文件夹，返回 [(器件名, 文件夹路径)]。"""
    base_dir = base_dir or P.CP_BASE_DIR
    out = []
    try:
        names = sorted(os.listdir(base_dir))
    except OSError:
        return out
    for name in names:
        path = os.path.join(base_dir, name)
        if not os.path.isdir(path):
            continue
        if name.startswith("__") or name.lower() == "thumbs.db":
            continue
        if name.upper().endswith("CP") and len(name) > 2:
            out.append((name[:-2], path))
    return out


def count_pages(data=None, base_dir=None):
    """本模块会生成多少页。"""
    return len(get_devices(base_dir))


def image_root(dev_folder):
    """优先取 xxxCP/CP 子目录，否则用 xxxCP 本身。"""
    sub = os.path.join(dev_folder, "CP")
    return sub if os.path.isdir(sub) else dev_folder


def profile_no(name):
    """1-1 / 1 -> 1；2-x -> 2 ... 非数字开头返回 None。"""
    m = re.match(r"\s*(\d+)", str(name))
    return int(m.group(1)) if m else None


def _list_entries(root):
    dirs, files = [], []
    try:
        names = sorted(os.listdir(root))
    except OSError:
        return dirs, files
    for name in names:
        if name.lower() == "thumbs.db":
            continue
        path = os.path.join(root, name)
        if os.path.isdir(path):
            dirs.append((name, path))
        else:
            files.append((name, path))
    return dirs, files


def _find_a_image(path):
    """文件夹中命名为 A 的图（排除 A- / AM）。"""
    if not os.path.isdir(path):
        return None
    cands = []
    for name in os.listdir(path):
        stem, ext = os.path.splitext(name)
        if stem == "A" and ext.lower() in IMG_EXTS:
            cands.append(os.path.join(path, name))
    return sorted(cands)[0] if cands else None


def collect_profiles(root):
    """{profile号: (容器路径, A图路径)}，按 profile 号升序。"""
    dirs, files = _list_entries(root)
    result = {}
    if dirs:
        for name, path in dirs:
            no = profile_no(name)
            if no is None or no in result:
                continue
            result[no] = (path, _find_a_image(path))
    else:
        # 平铺结构：文件名 1-x / 1 视为 profile1，A 图取根目录 A.*
        a_img = _find_a_image(root)
        nos = {profile_no(name) for name, _ in files}
        for no in nos:
            if no is not None:
                result[no] = (root, a_img)
    return {k: result[k] for k in sorted(result)}


# ============================================================
# SPEC / 实测值
# ============================================================
def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def hole_spec(data, device_name):
    """空洞率 SPEC（来自物料信息 SPEC 表的"空洞率"列）。"""
    v = data.get("物料信息", {}).get("spec", {}).get(device_name, {}).get("空洞率")
    return _num(v)


def hole_measured(data, device_name, profile_no_):
    """孔洞率实测值（来自"空洞率"sheet）；无数据返回 None。"""
    for key in ("空洞率", "孔洞率"):
        blk = data.get(key)
        if isinstance(blk, dict):
            sub = blk.get(profile_no_) or blk.get(f"Profile{profile_no_}")
            if isinstance(sub, dict) and device_name in sub:
                return _num(sub[device_name])
            if device_name in blk and not isinstance(blk[device_name], dict):
                return _num(blk[device_name])
    spec = data.get("物料信息", {}).get("spec", {}).get(device_name, {})
    for k in ("空洞率实测", "孔洞率实测"):
        if k in spec:
            return _num(spec[k])
    return None


def _pct(v):
    """单元格里已是百分数（0.6 = 0.6%），直接加 % 号。"""
    return "/" if v is None else f"{v:g}%"


# ============================================================
# 表格
# ============================================================
def _fill_table(shape, data, device_name, profiles):
    """填充表格：一行一个 profile。返回 [(profile号, A图路径, 是否满足SPEC)]。"""
    table = shape.table
    ref = U.get_ref_rPr(table)

    need = 1 + len(profiles)
    while len(table.rows) < need:
        U.add_table_row(table)
    if len(table.rows) > need:
        U.remove_table_rows(table, need)

    spec = hole_spec(data, device_name)
    # profile 名（来自"空洞率"sheet），缺省回退 Profile{n}
    names = data.get("空洞率_profiles", {}).get(device_name, {}) or {}

    results = []
    for ri, (no, (_, a_img)) in enumerate(profiles.items()):
        row = ri + 1
        meas = hole_measured(data, device_name, no)
        if spec is None or meas is None:
            row_concl, ok = "待定", None
        else:
            ok = meas <= spec
            row_concl = "符合" if ok else "不符合"
        vals = [str(names.get(no) or f"Profile{no}"), "",
                _pct(meas), _pct(spec), row_concl]
        for ci, v in enumerate(vals):
            if ci < len(table.columns):
                U.set_cell_text(table.cell(row, ci), v, ref)
        results.append((no, a_img, ok))

    # 防御性清空富余行
    for row in range(need, len(table.rows)):
        for ci in range(len(table.columns)):
            U.set_cell_text(table.cell(row, ci), "", ref)

    U.sync_shape_height(shape, table)
    return results


# ============================================================
# 结论
# ============================================================
def _page_conclusion(device_name, results, spec):
    if spec is None:
        return f"结论：{device_name}空洞率SPEC未定义，无法判定"
    oks = [ok for _, _, ok in results]
    if all(ok is None for ok in oks):
        return f"结论：{device_name}空洞率暂无实测数据，无法判定"
    if any(ok is False for ok in oks):
        return f"结论：{device_name}空洞率存在不符合SPEC的数据"
    if all(ok for ok in oks):
        return f"结论：{device_name}空洞率符合SPEC"
    return f"结论：{device_name}空洞率暂无实测数据，无法判定"


def _set_conclusion(slide, text, shape_bottom_emu=None):
    """写整页结论。shape_bottom_emu 给定则贴在表格下方，否则用默认 top。"""
    target = U.find_text_shape(slide, CONCL_PREFIX)
    if target is None:
        top = (Cm(P.CP_CONCL_DEFAULT_TOP_CM) if shape_bottom_emu is None
               else shape_bottom_emu + Cm(P.CP_CONCL_GAP_CM))
        target = slide.shapes.add_textbox(Cm(P.CP_CONCL_LEFT_CM), top,
                                          Cm(P.CP_CONCL_W_CM), Cm(P.CP_CONCL_H_CM))
    else:
        target.left = Cm(P.CP_CONCL_LEFT_CM)
        target.width = Cm(P.CP_CONCL_W_CM)
        target.height = Cm(P.CP_CONCL_H_CM)
        if shape_bottom_emu is not None:
            target.top = shape_bottom_emu + Cm(P.CP_CONCL_GAP_CM)
    U.write_text_box(target, text, FONT_LATIN, FONT_EA, FONT_SIZE)
    return text


# ============================================================
# 单页填充
# ============================================================
def display_path(path, base_dir):
    """日志展示用路径：base_dir 内显示相对路径，跨盘符显示绝对路径。"""
    try:
        return os.path.relpath(path, base_dir)
    except ValueError:
        return path


def fill_device_page(slide, device_name, dev_folder, data, base_dir):
    if not U.set_title(slide, TITLE_PREFIX, f"{TITLE_PREFIX} - {device_name}"):
        print(f"  [CP] 未找到标题文本框: {device_name}")

    root = image_root(dev_folder)
    profiles = collect_profiles(root)
    print(f"  [CP] {device_name}: profile {list(profiles.keys())} "
          f"(图片目录 {display_path(root, base_dir)})")

    shape = U.table_shape(slide)
    if shape is None:
        print(f"  [CP] {device_name} 页面无表格")
        results = []
    else:
        results = _fill_table(shape, data, device_name, profiles)
        for ri, (_, a_img, _) in enumerate(results):
            U.insert_picture_in_cell(slide, shape, ri + 1, COL_IMG, a_img)

    text = _page_conclusion(device_name, results, hole_spec(data, device_name))
    bottom = (shape.top + shape.height) if shape is not None else None
    print(f"  [CP] {device_name}: {_set_conclusion(slide, text, bottom)}")


def fill_cp_pages(prs, data, template_index=None, base_dir=None):
    """填充 CP 切片页：每个 <器件>CP 文件夹一页。

    页面复制已由主流程规划完成（模板切片页即第 1 个器件的页面），
    本函数只往 [template_index, template_index+count) 里填内容。返回填充页数。

    注：模板切片页上没有任何图片，因此直接复用它是安全的；
    若将来模板切片页带了示例图，需改为复制页再删除原页。
    """
    if template_index is None:
        template_index = P.TEMPLATE_PAGE_INDEX["cp"]

    devices = get_devices(base_dir)
    if not devices:
        print("[CP] 未发现 xxxCP 文件夹")
        return 0

    for i, (dev_name, folder) in enumerate(devices):
        fill_device_page(prs.slides[template_index + i],
                         dev_name, folder, data, base_dir or P.CP_BASE_DIR)

    print(f"[CP] 共生成 {len(devices)} 页CP切片: {[d for d, _ in devices]}")
    return len(devices)
