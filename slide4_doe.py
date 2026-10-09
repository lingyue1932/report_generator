# -*- coding: utf-8 -*-
"""第四页 DOE参数表填充"""
import copy

from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

FONT_EA = "宋体"
FONT_LATIN = "Times New Roman"
TABLE_FONT_SIZE_PT = 14

# 【新增】表格最小高度（行数少时防止表格过矮）
MIN_TABLE_HEIGHT = Cm(12)

# 【新增】Profile 列字体颜色；None 表示沿用模板原色
PROFILE_FONT_COLOR = "FFFFFF"

# 【新增】需要按 profile 组纵向合并的列（0-based）
#   0 = Profile；2 = Sample Size
MERGED_COLS = (0, 2)

# 与模板列顺序一致（模板列头含 Soler 拼写，保持不动）
PARAM_KEYS = [
    "Pre-heating Temp",
    "Ramp rate",
    "Solder Temp",
    "Press Time",
    "Solder Time",
    "Cooling Rate",
    "Cooling Temp",
    "Solder Force",
]


def _cell_val(v):
    """空格或/ 一律显示为 /"""
    if v is None:
        return "/"
    if isinstance(v, str):
        s = v.strip()
        if s == "" or s == "/":
            return "/"
        return s
    return str(v)


def _pretty_profile(name):
    """profile1 -> Profile 1；其余原样"""
    s = str(name).strip()
    low = s.lower().replace(" ", "")
    if low.startswith("profile") and low[7:].isdigit():
        return f"Profile {low[7:]}"
    return s


def _set_run_font(run, latin_font, ea_font, size_pt=None):
    run.font.name = latin_font
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn('a:latin'))
    if latin is None:
        latin = rPr.makeelement(qn('a:latin'), {'typeface': latin_font})
        rPr.append(latin)
    else:
        latin.set('typeface', latin_font)
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {'typeface': ea_font})
        latin.addnext(ea)
    else:
        ea.set('typeface', ea_font)
    cs = rPr.find(qn('a:cs'))
    if cs is None:
        cs = rPr.makeelement(qn('a:cs'), {'typeface': latin_font})
        ea.addnext(cs)
    else:
        cs.set('typeface', latin_font)


def _apply_font(shape, latin_font=FONT_LATIN, ea_font=FONT_EA,
                size_pt=TABLE_FONT_SIZE_PT):
    table = shape.table
    for row in table.rows:
        for cell in row.cells:
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    _set_run_font(run, latin_font, ea_font, size_pt)
            for latin in cell._tc.iter(qn('a:latin')):
                latin.set('typeface', latin_font)
            for ea in cell._tc.iter(qn('a:ea')):
                ea.set('typeface', ea_font)
            for cs in cell._tc.iter(qn('a:cs')):
                cs.set('typeface', latin_font)
            if size_pt is not None:
                sz = str(int(size_pt * 100))
                for def_rpr in cell._tc.iter(qn('a:defRPr')):
                    def_rpr.set('sz', sz)
                for end_rpr in cell._tc.iter(qn('a:endParaRPr')):
                    end_rpr.set('sz', sz)


def _set_cell_text(cell, text):
    """写单元格文本，保留原 run 的字符格式（颜色/加粗/字号等）。

    复用首个已存在的 run（继承其 rPr），只改文本；其余 run 删除。
    这样模板里预设的字体颜色（例如深底白字）不会丢。
    若原本没有任何 run，则保留段落 pPr，重新挂一个新 run。
    """
    tf = cell.text_frame
    if not tf.paragraphs:
        tf.add_paragraph()
    p = tf.paragraphs[0]

    for extra in list(tf.paragraphs[1:]):
        extra._p.getparent().remove(extra._p)

    runs = list(p.runs)
    if runs:
        keep = runs[0]
        keep.text = str(text)
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        pPr_tag = qn('a:pPr')
        for child in list(p._p):
            if child.tag == pPr_tag:
                continue
            p._p.remove(child)
        run = p.add_run()
        run.text = str(text)


def _set_col_font_color(table, col_idx, rgb_hex, start_row=0):
    """把指定列、从 start_row 起的单元格字体颜色设为指定色。"""
    if not rgb_hex:
        return
    rgb = RGBColor.from_string(rgb_hex)
    for ri in range(start_row, len(table.rows)):
        cell = table.cell(ri, col_idx)
        for para in cell.text_frame.paragraphs:
            for run in para.runs:
                run.font.color.rgb = rgb


def _add_row(table, ref_height=None):
    """复制最后一行追加一行，显式指定新行高度。"""
    rows = table._tbl.findall(qn('a:tr'))
    if not rows:
        return
    assert len(rows) >= 2, "模板表格至少需要表头 + 1 数据行，才能安全复制数据行"
    new_tr = copy.deepcopy(rows[-1])
    if ref_height is not None:
        new_tr.set('h', str(int(ref_height)))
    table._tbl.append(new_tr)


def _remove_extra_rows(table, keep):
    """删除第 keep 行之后的多余行"""
    trs = table._tbl.findall(qn('a:tr'))
    for tr in trs[keep:]:
        table._tbl.remove(tr)


def _sync_shape_size(shape, table, header_h=None, min_height=MIN_TABLE_HEIGHT):
    """同步表格尺寸：先统一行高，再让总高不低于 min_height。

    【关键】PPT 表格实际渲染高度 = 各行行高之和，只改 shape.height 不会
    把表格视觉撑开。因此当行高之和 < min_height 时，把差额均摊到数据行
    上，这样表格实际会变高。
    """
    n = len(table.rows)
    if n == 0:
        return
    if header_h is None:
        header_h = table.rows[0].height
    data_h = table.rows[1].height if n > 1 else header_h

    total = header_h + data_h * (n - 1)

    # 行高不足：均摊到数据行
    if min_height is not None and total < min_height and n > 1:
        data_h = (min_height - header_h) // (n - 1)
        total = header_h + data_h * (n - 1)

    table.rows[0].height = header_h
    for i in range(1, n):
        table.rows[i].height = data_h

    shape.height = total


def _clear_col_merge(table, col_idx):
    """清除指定列的合并属性，便于重新分组"""
    for tr in table._tbl.findall(qn('a:tr')):
        tcs = tr.findall(qn('a:tc'))
        if col_idx >= len(tcs):
            continue
        tc = tcs[col_idx]
        for attr in ('rowSpan', 'vMerge', 'hMerge', 'gridSpan'):
            if attr in tc.attrib:
                del tc.attrib[attr]


def _apply_col_merge(table, col_idx, group_sizes):
    """按组设置指定列纵向合并。group_sizes: 各组数据行数（不含表头）"""
    trs = table._tbl.findall(qn('a:tr'))
    ri = 1  # 跳过表头 row0
    for n in group_sizes:
        if n <= 0:
            continue
        tcs_head = trs[ri].findall(qn('a:tc'))
        if col_idx >= len(tcs_head):
            ri += n
            continue
        tc0 = tcs_head[col_idx]
        if n > 1:
            tc0.set('rowSpan', str(n))
        for j in range(1, n):
            tcj = trs[ri + j].findall(qn('a:tc'))[col_idx]
            tcj.set('vMerge', '1')
            for t in tcj.iter(qn('a:t')):
                t.text = ""
        ri += n


def _dev_params_empty(dev):
    """器件全部工艺参数是否为空（None / 空串 / 斜杠）"""
    for key in PARAM_KEYS:
        v = dev.get(key)
        if v is None:
            continue
        if isinstance(v, str):
            if v.strip() not in ("", "/"):
                return False
        else:
            return False
    return True


def build_doe_rows(data):
    """构建数据行（不含表头）。"""
    groups = data.get("物料信息", {}).get("doe_groups", [])
    out = []
    for g in groups:
        name = _pretty_profile(g.get("profile", ""))
        devices = [d for d in g.get("devices", []) if not _dev_params_empty(d)]
        if not devices:
            continue
        out.append((name, devices, g.get("sample_size")))
    return out


def fill_slide4(prs, data):
    """填充第四页DOE参数表"""
    if len(prs.slides) < 4:
        return

    slide = prs.slides[3]
    shapes = [s for s in slide.shapes if s.has_table]
    if not shapes:
        return
    shape = shapes[0]
    table = shape.table

    header_h = table.rows[0].height
    data_h = table.rows[1].height if len(table.rows) > 1 else header_h

    groups = build_doe_rows(data)
    total_data_rows = sum(len(devs) for _, devs, _ in groups)

    need = 1 + total_data_rows
    while len(table.rows) < need:
        _add_row(table, ref_height=data_h)
    if len(table.rows) > need:
        _remove_extra_rows(table, need)

    # 表头
    if len(table.columns) > 0:
        _set_cell_text(table.cell(0, 0), "Profile")

    # 清合并（顺序：先删/加行 → 再清合并 → 再写文本 → 最后重新合并）
    for ci in MERGED_COLS:
        _clear_col_merge(table, ci)

    r = 1
    group_sizes = []
    for profile_name, devices, sample_size in groups:
        n = len(devices)
        group_sizes.append(n)
        sample_s = _cell_val(sample_size)
        for di, dev in enumerate(devices):
            is_head = (di == 0)
            _set_cell_text(table.cell(r, 0), profile_name if is_head else "")
            _set_cell_text(table.cell(r, 1), _cell_val(dev.get("NO.")))
            _set_cell_text(table.cell(r, 2), sample_s if is_head else "")
            for pi, key in enumerate(PARAM_KEYS):
                _set_cell_text(table.cell(r, 3 + pi), _cell_val(dev.get(key)))
            r += 1

    while r < len(table.rows):
        for ci in range(len(table.columns)):
            _set_cell_text(table.cell(r, ci), "")
        r += 1

    # 重新合并
    for ci in MERGED_COLS:
        _apply_col_merge(table, ci, group_sizes)

    _apply_font(shape)

    # 强制 Profile 列数据单元格字体为白色（模板空单元格无可继承 run 时兜底）
    if PROFILE_FONT_COLOR:
        _set_col_font_color(table, 0, PROFILE_FONT_COLOR, start_row=1)

    # 同步表格尺寸：行高之和不足 MIN_TABLE_HEIGHT 时均摊到数据行，表格实际变高
    _sync_shape_size(shape, table, header_h=header_h)