# -*- coding: utf-8 -*-
"""可靠性页（模板第 6 页起）：每个器件一页。

页面内容：标题 + 居中 Die shear 可靠性图（33×11cm）+ 图表下方结论。
"""
import os

from pptx.util import Cm

import parameter as P
import pptx_utils as U
from data_reader import build_reliability_view

TITLE_PREFIX = "可靠性"
CONCL_PREFIX = "结论"

FONT_EA = P.FONT_EA_PPT
FONT_LATIN = P.FONT_LATIN_PPT
FONT_SIZE = P.FONT_SIZE_PPT_BODY

DIM_SHEAR = "剪切力"


def get_devices(data, view):
    """器件列表：先按物料信息 SPEC 顺序，其余按出现顺序追加。"""
    spec = data.get("物料信息", {}).get("spec", {})
    present, seen = [], set()
    for cond_data in view.values():
        for prof_data in cond_data.values():
            for dev_name in prof_data:
                if dev_name not in seen:
                    seen.add(dev_name)
                    present.append(dev_name)
    ordered = [d for d in spec if d in seen]
    ordered += [d for d in present if d not in ordered]
    return ordered


def count_pages(data):
    """本模块会生成多少页。"""
    return len(get_devices(data, build_reliability_view(data)))


def _dev_dims(dev_data):
    """该器件的有效维度（排除 SN / 统计）。"""
    return [k for k in dev_data if k not in ("SN", "统计")]


def _dim_values(dev_data, dim_name):
    """取指定维度的原始数值列表；dim_name 为空时取第一个有效维度。"""
    if dim_name is None:
        dims = _dev_dims(dev_data)
        if not dims:
            return None, []
        dim_name = dims[0]
    raw = dev_data.get(dim_name)
    if isinstance(raw, list):
        vals = [float(v) for v in raw
                if isinstance(v, (int, float)) and not isinstance(v, bool)]
        return dim_name, vals
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        return dim_name, [float(raw)]
    return dim_name, []


def device_all_pass(data, device_name, view):
    """该器件全部可靠性数据是否满足剪切力 SPEC。"""
    spec = data.get("物料信息", {}).get("spec", {})
    try:
        bound = float(spec.get(device_name, {}).get("剪切力"))
    except (TypeError, ValueError):
        return False

    checked = False
    for cond_data in view.values():
        for prof, prof_data in cond_data.items():
            if prof == "统计":
                continue
            dev_data = prof_data.get(device_name)
            if not isinstance(dev_data, dict):
                continue
            # 优先剪切力维度，否则取第一个有效维度
            dim_name = DIM_SHEAR if DIM_SHEAR in dev_data else None
            _dim_name, vals = _dim_values(dev_data, dim_name)
            for v in vals:
                checked = True
                if v < bound:
                    return False
    return checked


def _insert_chart(prs, slide, device_name):
    img = os.path.join(P.RELIABILITY_IMG_DIR, f"{device_name}_reliability.png")
    if not os.path.exists(img):
        print(f"  [REL] 图片不存在: {img}")
        return False
    left = (prs.slide_width - Cm(P.RELIABILITY_IMG_W_CM)) // 2
    slide.shapes.add_picture(
        img, left, Cm(P.RELIABILITY_IMG_TOP_CM),
        width=Cm(P.RELIABILITY_IMG_W_CM), height=Cm(P.RELIABILITY_IMG_H_CM))
    print(f"  [REL] {device_name} -> {os.path.basename(img)} "
          f"({P.RELIABILITY_IMG_W_CM}x{P.RELIABILITY_IMG_H_CM}cm)")
    return True


def _set_conclusion(slide, device_name, all_pass):
    if all_pass:
        text = f"结论：{device_name}在可靠性测试中Die shear均满足SPEC"
    else:
        text = f"结论：{device_name}在可靠性测试中Die shear存在不满足SPEC的数据"

    target = U.find_text_shape(slide, CONCL_PREFIX)
    if target is None:
        target = slide.shapes.add_textbox(
            Cm(P.RELIABILITY_CONCL_LEFT_CM), Cm(P.RELIABILITY_CONCL_TOP_CM),
            Cm(P.RELIABILITY_CONCL_W_CM), Cm(P.RELIABILITY_CONCL_H_CM))
    else:
        target.left = Cm(P.RELIABILITY_CONCL_LEFT_CM)
        target.top = Cm(P.RELIABILITY_CONCL_TOP_CM)
        target.width = Cm(P.RELIABILITY_CONCL_W_CM)
        target.height = Cm(P.RELIABILITY_CONCL_H_CM)
    U.write_text_box(target, text, FONT_LATIN, FONT_EA, FONT_SIZE)
    return text


def fill_device_page(prs, slide, device_name, data, view):
    if not U.set_title(slide, TITLE_PREFIX, f"{TITLE_PREFIX} - {device_name}"):
        print(f"  [REL] 未找到标题文本框: {device_name}")
    _insert_chart(prs, slide, device_name)
    text = _set_conclusion(slide, device_name,
                           device_all_pass(data, device_name, view))
    print(f"  [REL] {device_name}: {text}")


def fill_reliability_pages(prs, data, template_index=None):
    """填充可靠性页：每个器件一页。

    页面复制已由主流程完成，本函数只往
    [template_index, template_index+count) 里填内容。返回填充的页数。
    """
    if template_index is None:
        template_index = P.TEMPLATE_PAGE_INDEX["reliability"]

    view = build_reliability_view(data)
    devices = get_devices(data, view)
    if not devices:
        print("[REL] 无可靠性器件")
        return 0

    for i, dev_name in enumerate(devices):
        fill_device_page(prs, prs.slides[template_index + i],
                         dev_name, data, view)

    print(f"[REL] 共生成 {len(devices)} 页可靠性: {devices}")
    return len(devices)
