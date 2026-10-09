# -*- coding: utf-8 -*-
"""可靠性页：每个器件生成一页，居中插入可靠性图 + 结论
页面参考模板第六页
"""
import os

from pptx.util import Cm

import parameter as P
from slide_position import (_duplicate_slide, _reorder_slide,
                            _set_run_font, FONT_EA, FONT_LATIN)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

TEMPLATE_SLIDE_INDEX = 5  # 模板第六页（0-based）

FONT_SIZE = 14

# 图片尺寸 / 结论文本框坐标（cm）：统一在 parameter.py 配置

TITLE_PREFIX = "可靠性"


def build_all_data(data):
    """整理可靠性数据：{condition: {profile: {device: {dim: [values]}}}}"""
    all_data = {}
    for cond, (sheet_key, block_name) in P.RELIABILITY_MAP.items():
        if sheet_key not in data:
            continue
        sheet_data = data[sheet_key]
        if block_name not in sheet_data:
            continue
        all_data[cond] = {}
        for prof, content in sheet_data[block_name].items():
            if not isinstance(content, dict):
                continue
            all_data[cond][prof] = {
                dev_name: dev_data
                for dev_name, dev_data in content.items()
                if dev_name != "SN" and isinstance(dev_data, dict)
            }
    return all_data


def get_devices(data, all_data):
    """器件列表：先按物料信息SPEC顺序，其余按出现顺序追加"""
    spec = data.get("物料信息", {}).get("spec", {})
    present, seen = [], set()
    for cond_data in all_data.values():
        for prof_data in cond_data.values():
            for dev_name in prof_data:
                if dev_name not in seen:
                    seen.add(dev_name)
                    present.append(dev_name)
    ordered = [d for d in spec if d in seen]
    ordered += [d for d in present if d not in ordered]
    return ordered


def _dev_dims(dev_data):
    """该器件的有效维度（排除 SN / 统计）"""
    return [k for k in dev_data if k not in ("SN", "统计")]


def _dim_values(dev_data, dim_name):
    """取指定维度的原始数值列表；dim_name 为空时取第一个有效维度"""
    if dim_name is None:
        dims = _dev_dims(dev_data)
        if not dims:
            return None, []
        dim_name = dims[0]
    raw = dev_data.get(dim_name)
    if isinstance(raw, list):
        vals = [float(v) for v in raw
                if isinstance(v, (int, float)) and not isinstance(v, bool)]
        return dim_name, vals
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        return dim_name, [float(raw)]
    return dim_name, []


def device_all_pass(data, device_name, all_data):
    """该器件全部可靠性数据是否满足剪切力SPEC"""
    spec = data.get("物料信息", {}).get("spec", {})
    try:
        bound = float(spec.get(device_name, {}).get("剪切力"))
    except (TypeError, ValueError):
        return False

    checked = False
    for cond_data in all_data.values():
        for prof, prof_data in cond_data.items():
            if prof == "统计":
                continue
            dev_data = prof_data.get(device_name)
            if not isinstance(dev_data, dict):
                continue
            # 优先剪切力维度，否则取第一个有效维度
            dim_name = "剪切力" if "剪切力" in dev_data else None
            dim_name, vals = _dim_values(dev_data, dim_name)
            for v in vals:
                checked = True
                if v < bound:
                    return False
    return checked


def _set_title(slide, device_name):
    """标题：可靠性 - 器件名"""
    for shp in slide.shapes:
        if not shp.has_text_frame:
            continue
        txt = shp.text_frame.text.strip()
        if not txt.startswith(TITLE_PREFIX):
            continue
        tf = shp.text_frame
        title = f"{TITLE_PREFIX} - {device_name}"
        p = tf.paragraphs[0]
        if p.runs:
            p.runs[0].text = title
            for r in p.runs[1:]:
                r.text = ""
        else:
            run = p.add_run()
            run.text = title
        for extra in tf.paragraphs[1:]:
            for r in extra.runs:
                r.text = ""
        return True
    print(f"  [REL] 未找到标题文本框: {device_name}")
    return False


def _insert_image(prs, slide, device_name):
    img = os.path.join(P.RELIABILITY_IMG_DIR, f"{device_name}_reliability.png")
    if not os.path.exists(img):
        print(f"  [REL] 图片不存在: {img}")
        return False
    left = (prs.slide_width - Cm(P.RELIABILITY_IMG_W_CM)) // 2
    slide.shapes.add_picture(
        img, left, Cm(P.RELIABILITY_IMG_TOP_CM),
        width=Cm(P.RELIABILITY_IMG_W_CM), height=Cm(P.RELIABILITY_IMG_H_CM))
    print(f"  [REL] {device_name} -> {os.path.basename(img)} "
          f"({P.RELIABILITY_IMG_W_CM}x{P.RELIABILITY_IMG_H_CM}cm)")
    return True


def _set_conclusion(slide, device_name, all_pass):
    if all_pass:
        text = f"结论：{device_name}在可靠性测试中Die shear均满足SPEC"
    else:
        text = f"结论：{device_name}在可靠性测试中Die shear存在不满足SPEC的数据"

    target = None
    for shp in slide.shapes:
        if shp.has_text_frame and shp.text_frame.text.strip().startswith("结论"):
            target = shp
            break
    if target is None:
        target = slide.shapes.add_textbox(
            Cm(P.RELIABILITY_CONCL_LEFT_CM), Cm(P.RELIABILITY_CONCL_TOP_CM),
            Cm(P.RELIABILITY_CONCL_W_CM), Cm(P.RELIABILITY_CONCL_H_CM))
    else:
        target.left = Cm(P.RELIABILITY_CONCL_LEFT_CM)
        target.top = Cm(P.RELIABILITY_CONCL_TOP_CM)
        target.width = Cm(P.RELIABILITY_CONCL_W_CM)
        target.height = Cm(P.RELIABILITY_CONCL_H_CM)

    tf = target.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    _set_run_font(run, FONT_LATIN, FONT_EA, FONT_SIZE)
    return text


def _fill_device_page(prs, slide, device_name, data, all_data):
    _set_title(slide, device_name)
    _insert_image(prs, slide, device_name)
    all_pass = device_all_pass(data, device_name, all_data)
    text = _set_conclusion(slide, device_name, all_pass)
    print(f"  [REL] {device_name}: {text}")


def fill_reliability_pages(prs, data, template_index=TEMPLATE_SLIDE_INDEX):
    """生成可靠性页：每个器件一页
    - template_index 页作为第一个器件的页面，其余复制模板页插入其后
    - 图片居中，固定 33 x 11 cm
    - 图表下方写结论
    返回生成的页数
    """
    all_data = build_all_data(data)
    devices = get_devices(data, all_data)
    if not devices:
        print("[REL] 无可靠性器件")
        return 0

    sldIdLst = prs.slides._sldIdLst
    for i in range(1, len(devices)):
        _duplicate_slide(prs, template_index)
        _reorder_slide(prs, len(sldIdLst) - 1, template_index + i)

    for i, dev_name in enumerate(devices):
        slide = prs.slides[template_index + i]
        _fill_device_page(prs, slide, dev_name, data, all_data)

    print(f"[REL] 共生成 {len(devices)} 页可靠性: {devices}")
    return len(devices)
