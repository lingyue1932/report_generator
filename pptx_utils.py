# -*- coding: utf-8 -*-
"""PPT 操作基元层。

这些函数原先在 report_generator / slide4_doe / slide_position / slide_CP
里各有一份（或数份）几乎相同、或略有差异的实现：

    set_run_font       —— 3 份（report_generator、slide4_doe、slide_position）
    set_cell_text      —— 3 个变体（report_generator、slide4_doe、slide_CP）
    表格增删行          —— 2 套 API（_add_table_row/_remove_table_rows 对比
                          _add_row/_remove_extra_rows）
    _duplicate_slide   —— slide_position 独占，slide_reliability / slide_CP 借用
    _table_shape / _ref_rPr / _cell_rect —— slide_CP 与 slide_failure_mode 各一份

统一收拢到这里，各 page_* 模块只负责"这一页填什么内容"，
不再重复实现"怎么操作 pptx"。
"""
import copy

from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

import parameter as P


# ============================================================
# 字体
# ============================================================
def set_run_font(run, latin_font=None, ea_font=None, size_pt=P.FONT_SIZE_PPT_BODY):
    """设置 run 的字体：西文 latin_font、中文 ea_font（缺省同 latin）。

    latin_font / ea_font 为 None 时取 parameter 中的 PPT 默认值；
    size_pt 显式传 None 表示"不改动模板原有字号"。
    """
    if latin_font is None:
        latin_font = P.FONT_LATIN_PPT
    if ea_font is None:
        ea_font = latin_font

    run.font.name = latin_font
    if size_pt is not None:
        run.font.size = Pt(size_pt)

    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn('a:latin'))
    if latin is None:
        latin = rPr.makeelement(qn('a:latin'), {'typeface': latin_font})
        rPr.append(latin)
    else:
        latin.set('typeface', latin_font)

    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {'typeface': ea_font})
        latin.addnext(ea)
    else:
        ea.set('typeface', ea_font)

    cs = rPr.find(qn('a:cs'))
    if cs is None:
        cs = rPr.makeelement(qn('a:cs'), {'typeface': latin_font})
        ea.addnext(cs)
    else:
        cs.set('typeface', latin_font)


def apply_table_font(shape, latin_font=None, ea_font=None,
                     size_pt=P.FONT_SIZE_PPT_BODY):
    """统一整张表格的字体。

    除了逐个 run 设置，还要改单元格样式级的 a:latin / a:ea / a:cs
    以及 defRPr / endParaRPr 的 sz —— 否则空单元格与段落尾仍会沿用模板字体。
    """
    if latin_font is None:
        latin_font = P.FONT_LATIN_PPT
    if ea_font is None:
        ea_font = latin_font

    table = shape.table
    for row in table.rows:
        for cell in row.cells:
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    set_run_font(run, latin_font, ea_font, size_pt)
            for latin in cell._tc.iter(qn('a:latin')):
                latin.set('typeface', latin_font)
            for ea in cell._tc.iter(qn('a:ea')):
                ea.set('typeface', ea_font)
            for cs in cell._tc.iter(qn('a:cs')):
                cs.set('typeface', latin_font)
            if size_pt is not None:
                sz = str(int(size_pt * 100))
                for def_rpr in cell._tc.iter(qn('a:defRPr')):
                    def_rpr.set('sz', sz)
                for end_rpr in cell._tc.iter(qn('a:endParaRPr')):
                    end_rpr.set('sz', sz)


# ============================================================
# 单元格文本
# ============================================================
def set_cell_text(cell, text, ref_rPr=None):
    """写单元格文本，尽量保留模板原有的字符格式（颜色 / 加粗 / 字号）。

    策略：复用首个已存在的 run（继承其 rPr），只改文本，其余 run 与段落删除。
    这样模板里预设的格式（例如深底白字）不会丢。
    单元格原本没有任何 run 时，保留段落的 pPr（对齐等），新挂一个 run；
    若给了 ref_rPr，则把参照格式复制到新 run 上。
    """
    tf = cell.text_frame
    if not tf.paragraphs:
        tf.add_paragraph()
    p = tf.paragraphs[0]

    for extra in list(tf.paragraphs[1:]):
        extra._p.getparent().remove(extra._p)

    runs = list(p.runs)
    if runs:
        keep = runs[0]
        keep.text = str(text)
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        return

    pPr_tag = qn('a:pPr')
    for child in list(p._p):
        if child.tag == pPr_tag:
            continue
        p._p.remove(child)
    run = p.add_run()
    run.text = str(text)
    if ref_rPr is not None:
        old = run._r.find(qn('a:rPr'))
        new = copy.deepcopy(ref_rPr)
        if old is not None:
            run._r.replace(old, new)
        else:
            run._r.insert(0, new)


def get_ref_rPr(table):
    """取表格里第一个有文字的 run 的 rPr，作为无 run 单元格的格式参照。"""
    for row in table.rows:
        for cell in row.cells:
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    rPr = run._r.find(qn('a:rPr'))
                    if rPr is not None:
                        return copy.deepcopy(rPr)
    return None


# ============================================================
# 表格增删行 / 列
# ============================================================
def add_table_row(table, ref_height=None):
    """复制最后一行追加新行，保留格式。

    ref_height（EMU）不为 None 时显式写入新行高度。
    """
    rows = table._tbl.findall(qn('a:tr'))
    if not rows:
        return
    new_tr = copy.deepcopy(rows[-1])
    if ref_height is not None:
        new_tr.set('h', str(int(ref_height)))
    table._tbl.append(new_tr)


def remove_table_rows(table, keep):
    """删除第 keep 行之后的多余行（模板预置数据行多于实际数据时用）。"""
    trs = table._tbl.findall(qn('a:tr'))
    if keep < 1:
        keep = 1  # 至少保留表头
    for tr in trs[keep:]:
        table._tbl.remove(tr)


def add_table_col(table):
    """复制最后一列追加新列，保留格式，并均摊总宽度（总宽不变）。"""
    tbl = table._tbl
    grid = tbl.find(qn('a:tblGrid'))
    if grid is None:
        return
    cols = grid.findall(qn('a:gridCol'))
    if not cols:
        return
    total_w = sum(int(c.get('w', 0)) for c in cols)
    grid.append(copy.deepcopy(cols[-1]))
    cols = grid.findall(qn('a:gridCol'))
    n = len(cols)
    each = total_w // n if n else 0
    for c in cols:
        c.set('w', str(each))
    for tr in tbl.findall(qn('a:tr')):
        tcs = tr.findall(qn('a:tc'))
        if tcs:
            tr.append(copy.deepcopy(tcs[-1]))


def clear_col_merge(table, col_idx):
    """清除指定列的合并属性，便于按新分组重新合并。"""
    for tr in table._tbl.findall(qn('a:tr')):
        tcs = tr.findall(qn('a:tc'))
        if col_idx >= len(tcs):
            continue
        tc = tcs[col_idx]
        for attr in ('rowSpan', 'vMerge', 'hMerge', 'gridSpan'):
            if attr in tc.attrib:
                del tc.attrib[attr]


def apply_col_merge(table, col_idx, group_sizes):
    """按组对指定列做纵向合并。group_sizes: 各组数据行数（不含表头）。"""
    trs = table._tbl.findall(qn('a:tr'))
    ri = 1  # 跳过表头 row0
    for n in group_sizes:
        if n <= 0:
            continue
        tcs_head = trs[ri].findall(qn('a:tc'))
        if col_idx >= len(tcs_head):
            ri += n
            continue
        tc0 = tcs_head[col_idx]
        if n > 1:
            tc0.set('rowSpan', str(n))
        for j in range(1, n):
            tcj = trs[ri + j].findall(qn('a:tc'))[col_idx]
            tcj.set('vMerge', '1')
            for t in tcj.iter(qn('a:t')):
                t.text = ""
        ri += n


def fit_table_size(shape, width_cm=None, height_cm=None):
    """把表格缩放到指定宽×高（cm）；传 None 的维度沿用模板原尺寸。

    pptx 没有"整表缩放"API，只能逐列改宽、逐行改高；
    最后一列/行吸收累计舍入误差，保证总宽总高精确。
    """
    table = shape.table
    n_rows, n_cols = len(table.rows), len(table.columns)
    if not n_rows or not n_cols:
        return
    target_w = Cm(width_cm) if width_cm else shape.width
    target_h = Cm(height_cm) if height_cm else shape.height

    cur_w = sum(c.width for c in table.columns) or 1
    acc_w = 0
    for i, col in enumerate(table.columns):
        if i == n_cols - 1:
            col.width = target_w - acc_w
        else:
            w = int(col.width * target_w / cur_w)
            col.width = w
            acc_w += w

    row_h = target_h // n_rows
    acc_h = 0
    for i, row in enumerate(table.rows):
        if i == n_rows - 1:
            row.height = target_h - acc_h
        else:
            row.height = row_h
            acc_h += row_h

    shape.width = target_w
    shape.height = target_h


def sync_shape_height(shape, table):
    """行数变化后同步表格外框高度为实际行高之和，避免与表格错位。"""
    total = sum(r.height for r in table.rows)
    if total:
        shape.height = total
    return shape.height


# ============================================================
# 幻灯片复制 / 重排
# ============================================================
def duplicate_slide(prs, src_index):
    """复制幻灯片（深拷贝 spTree，含表格/文本，无图片关系），追加到末尾。

    注意：新页只带形状，不带图片关系；需要图片的页面须在填充阶段重新插入。
    """
    src = prs.slides[src_index]
    dest = prs.slides.add_slide(src.slide_layout)
    for shp in list(dest.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in src.shapes:
        dest.shapes._spTree.append(copy.deepcopy(shp._element))
    return dest


def reorder_slide(prs, from_index, to_index):
    """把第 from_index 页移到 to_index 位置。"""
    sldIdLst = prs.slides._sldIdLst
    el = list(sldIdLst)[from_index]
    sldIdLst.remove(el)
    sldIdLst.insert(to_index, el)


def drop_slide(prs, index):
    """删除第 index 页（含其部件关系）。"""
    sldIdLst = prs.slides._sldIdLst
    el = list(sldIdLst)[index]
    prs.part.drop_rel(el.get(qn('r:id')))
    sldIdLst.remove(el)


def expand_pages(prs, template_index, count):
    """把 template_index 页复制 count-1 份，紧随其后排开。

    调用后 prs.slides[template_index : template_index+count] 即为该组的全部页面，
    页面总数增加 count-1。count <= 0 时不改动。
    """
    if count <= 0:
        return 0
    sldIdLst = prs.slides._sldIdLst
    for i in range(1, count):
        duplicate_slide(prs, template_index)
        reorder_slide(prs, len(sldIdLst) - 1, template_index + i)
    return count


# ============================================================
# 形状查找 / 定位
# ============================================================
def find_text_shape(slide, prefix):
    """找文本框文本以 prefix 开头的第一个形状（strip 后比较）。"""
    for shp in slide.shapes:
        if shp.has_text_frame and shp.text_frame.text.strip().startswith(prefix):
            return shp
    return None


def find_shape_containing(slide, text):
    """找文本框文本包含 text 的第一个形状。"""
    for shp in slide.shapes:
        if shp.has_text_frame and text in shp.text_frame.text:
            return shp
    return None


def table_shape(slide):
    """取页面上第一张表格，没有则 None。"""
    for shp in slide.shapes:
        if shp.has_table:
            return shp
    return None


def set_title(slide, prefix, title):
    """把以 prefix 开头的标题文本框改写为 title，保留原格式。返回是否找到。"""
    shp = find_text_shape(slide, prefix)
    if shp is None:
        return False
    tf = shp.text_frame
    p = tf.paragraphs[0]
    if p.runs:
        p.runs[0].text = title
        for r in p.runs[1:]:
            r.text = ""
    else:
        run = p.add_run()
        run.text = title
    for extra in tf.paragraphs[1:]:
        for r in extra.runs:
            r.text = ""
    return True


def write_text_box(shape, text, latin_font=None, ea_font=None,
                   size_pt=P.FONT_SIZE_PPT_BODY):
    """清空文本框写入一行文字，并统一字体。"""
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run_font(run, latin_font, ea_font, size_pt)
    return text


# ============================================================
# 图片插入
# ============================================================
def cell_rect(shape, table, row, col):
    """返回单元格相对幻灯片的 (left, top, width, height)，单位 EMU。"""
    x = shape.left + sum(c.width for i, c in enumerate(table.columns) if i < col)
    y = shape.top + sum(r.height for i, r in enumerate(table.rows) if i < row)
    return x, y, table.columns[col].width, table.rows[row].height


def insert_picture_in_cell(slide, shape, row, col, img_path,
                           height_cm=P.CP_IMG_HEIGHT_CM):
    """按单元格位置插图：高度固定，宽度等比例，单元格内居中。"""
    import os

    from PIL import Image

    if not img_path or not os.path.exists(img_path):
        print(f"  [IMG] 图片不存在: {img_path}")
        return False
    try:
        with Image.open(img_path) as im:
            w, h = im.size
    except Exception as e:
        print(f"  [IMG] 图片无法读取 {img_path}: {e}")
        return False
    if not w or not h:
        return False

    img_h = Cm(height_cm)
    img_w = int(img_h * w / h)

    x, y, cw, rh = cell_rect(shape, shape.table, row, col)
    left = x + max(0, (cw - img_w) // 2)
    top = y + max(0, (rh - img_h) // 2)
    slide.shapes.add_picture(img_path, left, top, width=img_w, height=img_h)
    print(f"  [IMG] {os.path.basename(img_path)} ({img_w / 360000:.2f}x{height_cm}cm)")
    return True
