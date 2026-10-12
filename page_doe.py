# -*- coding: utf-8 -*-
"""第四页：DOE 参数表。

表结构：列 0=Profile，列 1=NO.，列 2=Sample size，列 3.. = 8 项工艺参数。
Profile / Sample size 两列按 profile 组纵向合并。

注意：本页模板历史上用"宋体"，其余各页用"微软雅黑"。当前保持原样，
如需统一，改 parameter.FONT_EA_PPT_DOE_TABLE 即可。
"""
from pptx.dml.color import RGBColor
from pptx.util import Cm

import parameter as P
import pptx_utils as U

FONT_EA = P.FONT_EA_PPT_DOE_TABLE
FONT_LATIN = P.FONT_LATIN_PPT
FONT_SIZE_PT = P.FONT_SIZE_PPT_BODY

# 表格最小高度（行数少时防止表格过矮）
MIN_TABLE_HEIGHT = Cm(12)

# Profile 列字体颜色；None 表示沿用模板原色
PROFILE_FONT_COLOR = "FFFFFF"

# 需要按 profile 组纵向合并的列（0-based）：0=Profile，2=Sample size
MERGED_COLS = (0, 2)

# 与模板列顺序一致（模板列头含 Soler 拼写，保持不动）
PARAM_KEYS = [
    "Pre-heating Temp",
    "Ramp rate",
    "Solder Temp",
    "Press Time",
    "Solder Time",
    "Cooling Rate",
    "Cooling Temp",
    "Solder Force",
]


def _cell_val(v):
    """空格或 / 一律显示为 /"""
    if v is None:
        return "/"
    if isinstance(v, str):
        s = v.strip()
        return "/" if s in ("", "/") else s
    return str(v)


def _pretty_profile(name):
    """profile1 -> Profile 1；其余原样"""
    s = str(name).strip()
    low = s.lower().replace(" ", "")
    if low.startswith("profile") and low[7:].isdigit():
        return f"Profile {low[7:]}"
    return s


def _set_col_font_color(table, col_idx, rgb_hex, start_row=0):
    """把指定列、从 start_row 起的单元格字体颜色设为指定色。"""
    if not rgb_hex:
        return
    rgb = RGBColor.from_string(rgb_hex)
    for ri in range(start_row, len(table.rows)):
        cell = table.cell(ri, col_idx)
        for para in cell.text_frame.paragraphs:
            for run in para.runs:
                run.font.color.rgb = rgb


def _sync_shape_size(shape, table, header_h=None, min_height=MIN_TABLE_HEIGHT):
    """同步表格尺寸：先统一行高，再让总高不低于 min_height。

    PPT 表格实际渲染高度 = 各行行高之和，只改 shape.height 不会把表格
    视觉撑开。因此当行高之和 < min_height 时，把差额均摊到数据行上。
    """
    n = len(table.rows)
    if n == 0:
        return
    if header_h is None:
        header_h = table.rows[0].height
    data_h = table.rows[1].height if n > 1 else header_h
    total = header_h + data_h * (n - 1)

    if min_height is not None and total < min_height and n > 1:
        data_h = (min_height - header_h) // (n - 1)
        total = header_h + data_h * (n - 1)

    table.rows[0].height = header_h
    for i in range(1, n):
        table.rows[i].height = data_h
    shape.height = total


def _dev_params_empty(dev):
    """器件全部工艺参数是否为空（None / 空串 / 斜杠）"""
    for key in PARAM_KEYS:
        v = dev.get(key)
        if v is None:
            continue
        if isinstance(v, str):
            if v.strip() not in ("", "/"):
                return False
        else:
            return False
    return True


def build_rows(data):
    """构建数据行（不含表头）：[(profile 显示名, [器件 dict...], sample_size)]"""
    groups = data.get("物料信息", {}).get("doe_groups", [])
    out = []
    for g in groups:
        devices = [d for d in g.get("devices", []) if not _dev_params_empty(d)]
        if not devices:
            continue
        out.append((_pretty_profile(g.get("profile", "")),
                    devices, g.get("sample_size")))
    return out


def fill_doe(prs, data):
    """填充第四页 DOE 参数表。返回数据行数。"""
    idx = P.TEMPLATE_PAGE_INDEX["doe"]
    if len(prs.slides) <= idx:
        return 0

    slide = prs.slides[idx]
    shape = U.table_shape(slide)
    if shape is None:
        return 0
    table = shape.table

    header_h = table.rows[0].height
    data_h = table.rows[1].height if len(table.rows) > 1 else header_h

    groups = build_rows(data)
    need = 1 + sum(len(devs) for _, devs, _ in groups)

    while len(table.rows) < need:
        U.add_table_row(table, ref_height=data_h)
    if len(table.rows) > need:
        U.remove_table_rows(table, need)

    # 表头
    if len(table.columns) > 0:
        U.set_cell_text(table.cell(0, 0), "Profile")

    # 顺序：先增删行 → 再清合并 → 再写文本 → 最后重新合并
    for ci in MERGED_COLS:
        U.clear_col_merge(table, ci)

    r = 1
    group_sizes = []
    for profile_name, devices, sample_size in groups:
        group_sizes.append(len(devices))
        sample_s = _cell_val(sample_size)
        for di, dev in enumerate(devices):
            is_head = (di == 0)
            U.set_cell_text(table.cell(r, 0), profile_name if is_head else "")
            U.set_cell_text(table.cell(r, 1), _cell_val(dev.get("NO.")))
            U.set_cell_text(table.cell(r, 2), sample_s if is_head else "")
            for pi, key in enumerate(PARAM_KEYS):
                U.set_cell_text(table.cell(r, 3 + pi), _cell_val(dev.get(key)))
            r += 1

    while r < len(table.rows):
        for ci in range(len(table.columns)):
            U.set_cell_text(table.cell(r, ci), "")
        r += 1

    for ci in MERGED_COLS:
        U.apply_col_merge(table, ci, group_sizes)

    U.apply_table_font(shape, FONT_LATIN, FONT_EA, FONT_SIZE_PT)

    # 强制 Profile 列数据单元格字体为白色（模板空单元格无可继承 run 时兜底）
    if PROFILE_FONT_COLOR:
        _set_col_font_color(table, 0, PROFILE_FONT_COLOR, start_row=1)

    _sync_shape_size(shape, table, header_h=header_h)
    return sum(group_sizes)
